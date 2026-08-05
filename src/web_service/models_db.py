"""SQLAlchemy 数据模型"""
import json
import logging
from datetime import datetime

logger = logging.getLogger(__name__)

from sqlalchemy import (
    Boolean,
    Column,
    Float,
    ForeignKey,
    Integer,
    String,
    Text,
    create_engine,
)
from sqlalchemy.orm import DeclarativeBase, Session, sessionmaker, relationship

from .config import DATABASE_URL


class Base(DeclarativeBase):
    pass


class User(Base):
    __tablename__ = "users"

    id = Column(Integer, primary_key=True)
    username = Column(String, unique=True, nullable=False)
    hashed_password = Column(String, nullable=False)
    role = Column(String, default="user")  # 'user' or 'admin'
    created_at = Column(String, default=lambda: datetime.now().isoformat())


class Task(Base):
    __tablename__ = "tasks"

    id = Column(Integer, primary_key=True)
    user_id = Column(Integer, ForeignKey("users.id"), nullable=False)
    batch_name = Column(String, nullable=False)
    status = Column(String, default="pending")
    file_count = Column(Integer, default=0)
    output_dir = Column(String)
    error_message = Column(Text)
    created_at = Column(String, default=lambda: datetime.now().isoformat())
    completed_at = Column(String)


class TaskFile(Base):
    __tablename__ = "task_files"

    id = Column(Integer, primary_key=True)
    task_id = Column(Integer, ForeignKey("tasks.id"), nullable=False)
    filename = Column(String, nullable=False)
    file_type = Column(String)
    parsed = Column(Boolean, default=False)
    md_path = Column(String)


class AuditRule(Base):
    __tablename__ = "audit_rules"

    id = Column(Integer, primary_key=True)
    category = Column(String, nullable=False)
    rule_name = Column(String, nullable=False)
    clause = Column(String, nullable=False)
    level = Column(String, nullable=False)
    description = Column(Text)
    check_expression = Column(Text)
    source_document = Column(String)
    enabled = Column(Boolean, default=True)
    created_by = Column(Integer, ForeignKey("users.id"))
    created_at = Column(String, default=lambda: datetime.now().isoformat())
    updated_at = Column(String, default=lambda: datetime.now().isoformat())


class AuditRuleHistory(Base):
    __tablename__ = "audit_rule_history"

    id = Column(Integer, primary_key=True)
    rule_id = Column(Integer, ForeignKey("audit_rules.id"), nullable=False)
    action = Column(String, nullable=False)
    old_value = Column(Text)
    new_value = Column(Text)
    changed_by = Column(Integer, ForeignKey("users.id"))
    changed_at = Column(String, default=lambda: datetime.now().isoformat())


class DocumentVersion(Base):
    __tablename__ = "document_versions"

    id = Column(Integer, primary_key=True)
    document_name = Column(String, nullable=False)
    file_path = Column(String, nullable=False)
    last_mtime = Column(Float, default=0.0)
    version_tag = Column(String)
    parsed_at = Column(String, default=lambda: datetime.now().isoformat())
    sync_status = Column(String, default="synced")
    sync_diff = Column(Text)
    applied_by = Column(Integer, ForeignKey("users.id"))
    applied_at = Column(String)


class SyncHistory(Base):
    __tablename__ = "sync_history"

    id = Column(Integer, primary_key=True)
    trigger_type = Column(String, nullable=False)
    document_name = Column(String, nullable=False)
    rules_added = Column(Integer, default=0)
    rules_modified = Column(Integer, default=0)
    rules_deleted = Column(Integer, default=0)
    diff_report = Column(Text)
    status = Column(String, default="pending")
    operated_by = Column(Integer, ForeignKey("users.id"))
    operated_at = Column(String, default=lambda: datetime.now().isoformat())


class ReceptionStandard(Base):
    """接待标准存储表 — 管理员可修改"""

    __tablename__ = "reception_standards"

    id = Column(Integer, primary_key=True)
    category = Column(String, nullable=False)  # "external" or "internal"
    sub_category = Column(String, nullable=False)  # "外事/商务" or "其他公务"
    personnel_level = Column(String, nullable=False)  # "省管中层" / "市管中层" / "其他人员"
    meal_limit = Column(Float, default=0.0)  # 用餐人均上限
    souvenir_limit = Column(Float, nullable=True)  # 纪念品上限，NULL表示不得赠送
    baijiu_limit = Column(Float, default=0.0)  # 白酒上限
    red_wine_limit = Column(Float, default=0.0)  # 红酒上限
    enabled = Column(Boolean, default=True)
    created_at = Column(String, default=lambda: datetime.now().isoformat())
    updated_at = Column(String, default=lambda: datetime.now().isoformat())
    updated_by = Column(Integer, ForeignKey("users.id"))


class ReceptionStandardHistory(Base):
    """接待标准修改历史"""

    __tablename__ = "reception_standard_history"

    id = Column(Integer, primary_key=True)
    standard_id = Column(Integer, ForeignKey("reception_standards.id"), nullable=False)
    action = Column(String, nullable=False)  # "create" / "update" / "delete"
    old_values = Column(Text)  # JSON
    new_values = Column(Text)  # JSON
    changed_by = Column(Integer, ForeignKey("users.id"))
    changed_at = Column(String, default=lambda: datetime.now().isoformat())


class ReceptionTemplate(Base):
    """场景接待模板存储表"""

    __tablename__ = "reception_templates"

    id = Column(Integer, primary_key=True)
    template_id = Column(String, unique=True, nullable=False)  # unique key
    title = Column(String, nullable=False)
    icon = Column(String)
    icon_class = Column(String)
    reception_type = Column(String, nullable=False)
    category = Column(String)
    example_data = Column(Text)  # JSON
    notes = Column(Text)
    required_docs = Column(Text)
    enabled = Column(Boolean, default=True)
    created_at = Column(String, default=lambda: datetime.now().isoformat())
    updated_at = Column(String, default=lambda: datetime.now().isoformat())
    updated_by = Column(Integer, ForeignKey("users.id"))


class ReceptionRecord(Base):
    """招待费提交记录 — 由 card7-fill.html 表单提交"""

    __tablename__ = "reception_records"

    id = Column(Integer, primary_key=True)
    serial_number = Column(String, unique=True, nullable=False)  # 提交流水号，如 ZDF-20260724-0001
    scenario = Column(String, nullable=False)  # A/B/C/D/E/F
    scenario_name = Column(String, nullable=False)
    submitter = Column(String, default="匿名")
    date = Column(String)  # 招待日期
    location = Column(String)
    org = Column(String)  # 招待对象单位
    guests = Column(Integer)  # 招待对象人数
    accompany_people = Column(String)  # 陪同人员
    accompany_count = Column(Integer)  # 陪同人数
    per_capita = Column(Float)  # 人均消费
    total_amount = Column(Float)  # 总金额
    baijiu_price = Column(String)  # 白酒单价
    wine_price = Column(String)  # 葡萄酒单价
    gift_per = Column(String)  # 礼品单价
    rooms = Column(String)  # 房间数
    room_price = Column(String)  # 房价
    reason = Column(Text)  # 事由
    nationality = Column(String)  # 外事招待-国籍
    fee = Column(String)  # 工作餐-费用
    status = Column(String, default="pending")  # pending / approved / rejected
    ocr_status = Column(String, default="pending")  # pending / processing / completed / skipped
    review_status = Column(String, default="pending")  # pending / reviewing / completed
    ocr_file_count = Column(Integer, default=0)  # 已上传OCR文件数
    review_report = Column(Text)  # 审核报告内容(Markdown)
    created_at = Column(String, default=lambda: datetime.now().isoformat())
    updated_at = Column(String, default=lambda: datetime.now().isoformat())


# --- Engine & Session ---

engine = create_engine(DATABASE_URL, connect_args={"check_same_thread": False})
SessionLocal = sessionmaker(bind=engine)


def init_db():
    Base.metadata.create_all(bind=engine)
    _migrate_add_columns()
    _migrate_slot_names()
    _seed_default_rules()
    _ensure_standards_seeded()
    _ensure_templates_seeded()
    _ensure_extraction_rules_seeded()
    _migrate_internal_units()
    _seed_internal_units()


def _migrate_add_columns():
    """增量迁移：为已有表补加缺失列（SQLite 不自动 ALTER）"""
    import sqlite3
    db_path = DATABASE_URL.replace("sqlite:///", "")
    conn = sqlite3.connect(db_path)
    try:
        cur = conn.cursor()
        # 获取所有表的列
        tables_to_check = {
            "ocr_records": [
                ("ocr_blocks", "TEXT"),
                ("ocr_lines", "TEXT"),
                ("img_width", "INTEGER", "0"),
                ("img_height", "INTEGER", "0"),
                ("engine", "TEXT"),
                ("original_file_path", "TEXT"),
                ("dedup_key", "TEXT"),
            ],
            "reception_records": [
                ("ocr_status", "TEXT"),
                ("review_status", "TEXT"),
                ("ocr_file_count", "INTEGER"),
                ("review_report", "TEXT"),
            ],
        }
        for table, columns in tables_to_check.items():
            cur.execute(f"PRAGMA table_info({table})")
            existing = {row[1] for row in cur.fetchall()}
            for col_def in columns:
                col_name = col_def[0]
                col_type = col_def[1]
                default = col_def[2] if len(col_def) > 2 else None
                if col_name not in existing:
                    try:
                        ddl = f"ALTER TABLE {table} ADD COLUMN {col_name} {col_type}"
                        if default is not None:
                            ddl += f" DEFAULT {default}"
                        cur.execute(ddl)
                        logger.info(f"Migrated: added column {col_name} to {table}")
                    except sqlite3.OperationalError:
                        pass  # 忽略已存在的列
        conn.commit()
    except Exception as e:
        logger.warning(f"Migration skipped or failed: {e}")
    finally:
        conn.close()


def _migrate_slot_names():
    """修正 OCR 记录中的 slot_name — 将前端旧版不规范的名称映射到后端标准名称"""
    import sqlite3
    db_path = DATABASE_URL.replace("sqlite:///", "")
    conn = sqlite3.connect(db_path)
    try:
        cur = conn.cursor()

        # 旧前端 slotNameMap → 后端 SLOTS 标准 name 的映射
        slot_mapping = {
            "发票(XML)": "发票XML",
            "电子发票(PDF)": "发票PDF",
            "招待单位经营状态": "经营状态",
            "支付流水证明": "支付流水",
        }

        # 检查是否需要迁移
        needs_migration = False
        for old_name in slot_mapping:
            cur.execute("SELECT COUNT(*) FROM ocr_records WHERE slot_name = ?", (old_name,))
            if cur.fetchone()[0] > 0:
                needs_migration = True
                break

        if not needs_migration:
            return

        tables_to_fix = ["ocr_records", "extracted_fields"]
        for table in tables_to_fix:
            for old_name, new_name in slot_mapping.items():
                cur.execute(f"UPDATE {table} SET slot_name = ? WHERE slot_name = ?", (new_name, old_name))
                affected = cur.rowcount
                if affected > 0:
                    logger.info(f"Migrated: {table} slot_name '{old_name}' → '{new_name}' ({affected} rows)")

        conn.commit()
    except Exception as e:
        logger.warning(f"Slot name migration failed: {e}")
    finally:
        conn.close()


def _ensure_standards_seeded():
    """确保标准数据已注入（独立于规则注入）"""
    db = SessionLocal()
    try:
        if db.query(ReceptionStandard).first():
            return
        _seed_default_standards(db)
        db.commit()
    except Exception:
        db.rollback()
    finally:
        db.close()


def _ensure_templates_seeded():
    """确保模板数据已注入（独立于规则注入）"""
    db = SessionLocal()
    try:
        if db.query(ReceptionTemplate).first():
            return
        _seed_default_templates(db)
        db.commit()
    except Exception:
        db.rollback()
    finally:
        db.close()


def _ensure_extraction_rules_seeded():
    """确保字段提取规则已注入"""
    db = SessionLocal()
    try:
        if db.query(ExtractionRule).first():
            return
        _seed_default_extraction_rules(db)
        db.commit()
    except Exception:
        db.rollback()
    finally:
        db.close()


def _seed_default_extraction_rules(db: Session):
    """注入默认字段提取规则"""
    rules = [
        {
            "slot_name": "审批单",
            "rule_name": "approval_extraction",
            "fields": [
                {"name": "招待日期", "keywords": ["招待日期", "日期"], "type": "date"},
                {"name": "招待人数", "keywords": ["招待人数", "对象人数", "招待对象人数"], "type": "number"},
                {"name": "陪同人数", "keywords": ["陪同人数", "陪同"], "type": "number"},
                {"name": "人均费用", "keywords": ["人均费用", "人均", "人均消费"], "type": "number"},
                {"name": "招待金额", "keywords": ["招待金额", "金额", "预算金额"], "type": "number"},
                {"name": "招待类型", "keywords": ["招待类型", "类型", "业务类型"], "type": "text"},
                {"name": "招待对象", "keywords": ["招待对象", "对象单位", "招待单位", "对方单位"], "type": "text"},
                {"name": "陪同人员", "keywords": ["陪同人员", "陪同"], "type": "text"},
                {"name": "事由", "keywords": ["事由", "事由及内容"], "type": "text"},
            ],
        },
        {
            "slot_name": "申请单",
            "rule_name": "application_extraction",
            "fields": [
                {"name": "申请日期", "keywords": ["申请日期", "日期"], "type": "date"},
                {"name": "事由", "keywords": ["事由", "申请事由", "申请内容"], "type": "text"},
                {"name": "预计费用", "keywords": ["预计费用", "预算", "金额"], "type": "number"},
                {"name": "参加人员", "keywords": ["参加人员", "人员"], "type": "text"},
            ],
        },
        {
            "slot_name": "报账单",
            "rule_name": "reimbursement_extraction",
            "fields": [
                {"name": "发票金额", "keywords": ["发票金额", "金额"], "type": "number"},
                {"name": "税额", "keywords": ["税额"], "type": "number"},
                {"name": "合计金额", "keywords": ["合计", "合计金额", "报销金额"], "type": "number"},
                {"name": "报销日期", "keywords": ["报销日期", "日期"], "type": "date"},
            ],
        },
        {
            "slot_name": "发票XML",
            "rule_name": "invoice_extraction",
            "fields": [
                {"name": "价税合计大写", "keywords": ["价税合计（大写)", "价税合计大写", "合计大写"], "type": "text"},
                {"name": "价税合计小写", "keywords": ["价税合计（小写)", "（小写）", "小写"], "type": "number"},
                {"name": "开票日期", "keywords": ["开票日期", "日期"], "type": "date"},
                {"name": "购买方", "keywords": ["购买方", "买方"], "type": "text"},
                {"name": "销售方", "keywords": ["销售方", "卖方"], "type": "text"},
                {"name": "发票号码", "keywords": ["发票号码", "号码"], "type": "text"},
            ],
        },
        {
            "slot_name": "发票PDF",
            "rule_name": "invoice_extraction",
            "fields": [
                {"name": "价税合计大写", "keywords": ["价税合计（大写)", "价税合计大写", "合计大写"], "type": "text"},
                {"name": "价税合计小写", "keywords": ["价税合计（小写)", "（小写）", "小写"], "type": "number"},
                {"name": "开票日期", "keywords": ["开票日期", "日期"], "type": "date"},
                {"name": "购买方", "keywords": ["购买方", "买方"], "type": "text"},
                {"name": "销售方", "keywords": ["销售方", "卖方"], "type": "text"},
                {"name": "发票号码", "keywords": ["发票号码", "号码"], "type": "text"},
            ],
        },
        {
            "slot_name": "经营状态",
            "rule_name": "business_status_extraction",
            "fields": [
                {"name": "单位名称", "keywords": ["单位名称", "企业名称"], "type": "text"},
                {"name": "经营状态", "keywords": ["经营状态", "状态"], "type": "text"},
            ],
        },
        {
            "slot_name": "支付凭证",
            "rule_name": "payment_extraction",
            "fields": [
                {"name": "交易单号", "keywords": ["交易单号", "单号"], "type": "text"},
                {"name": "商户单号", "keywords": ["商户单号"], "type": "text"},
                {"name": "金额", "keywords": ["金额", "账单金额"], "type": "number"},
                {"name": "商户全称", "keywords": ["商户全称", "商户名称", "商户"], "type": "text"},
                {"name": "支付时间", "keywords": ["支付时间", "时间"], "type": "date"},
            ],
        },
        {
            "slot_name": "支付流水",
            "rule_name": "payment_flow_extraction",
            "fields": [
                {"name": "交易单号", "keywords": ["交易单号", "单号"], "type": "text"},
                {"name": "交易对方", "keywords": ["交易对方", "对方", "收款方"], "type": "text"},
                {"name": "金额", "keywords": ["金额", "金额(元)"], "type": "number"},
                {"name": "交易时间", "keywords": ["交易时间", "时间"], "type": "date"},
            ],
        },
        {
            "slot_name": "活动函件",
            "rule_name": "event_extraction",
            "fields": [
                {"name": "活动名称", "keywords": ["活动名称", "活动", "会议名称", "函件主题"], "type": "text"},
                {"name": "交流时间", "keywords": ["交流时间", "活动时间", "时间"], "type": "date"},
                {"name": "参加人员", "keywords": ["参加人员", "参会人员", "人员"], "type": "text"},
            ],
        },
    ]
    for rule in rules:
        db.add(ExtractionRule(
            slot_name=rule["slot_name"],
            rule_name=rule["rule_name"],
            fields_json=json.dumps(rule["fields"], ensure_ascii=False),
            enabled=True,
        ))


def _seed_default_rules():
    """首次启动时从制度文档解析规则到数据库"""
    # Lazy import to avoid circular deps
    from .config import RULES_DIR
    from .rule_parser import parse_all_documents

    import json
    import os

    db = SessionLocal()
    try:
        if db.query(AuditRule).first():
            return

        rules = parse_all_documents()
        for rule_data in rules:
            rule = AuditRule(
                category=rule_data.get("category", ""),
                rule_name=rule_data.get("rule_name", ""),
                clause=rule_data.get("clause", ""),
                level=rule_data.get("level", "中"),
                description=rule_data.get("description", ""),
                check_expression=rule_data.get("check_expression", ""),
                source_document=rule_data.get("source_document", ""),
                enabled=True,
            )
            db.add(rule)

        for fpath in RULES_DIR.glob("*.docx"):
            dv = DocumentVersion(
                document_name=fpath.stem,
                file_path=str(fpath),
                last_mtime=os.path.getmtime(fpath),
                sync_status="synced",
            )
            db.add(dv)

        db.commit()
    except Exception:
        db.rollback()
    finally:
        db.close()


def _seed_default_standards(db: Session):
    """从 defaults 注入默认接待标准"""
    defaults = [
        # 对外
        {"category": "external", "sub_category": "外事/商务", "personnel_level": "省管中层", "meal_limit": 400.0, "souvenir_limit": 400.0, "baijiu_limit": 400.0, "red_wine_limit": 400.0},
        {"category": "external", "sub_category": "其他公务", "personnel_level": "省管中层", "meal_limit": 200.0, "souvenir_limit": None, "baijiu_limit": 400.0, "red_wine_limit": 400.0},
        {"category": "external", "sub_category": "外事/商务", "personnel_level": "市管中层", "meal_limit": 300.0, "souvenir_limit": 300.0, "baijiu_limit": 300.0, "red_wine_limit": 300.0},
        {"category": "external", "sub_category": "其他公务", "personnel_level": "市管中层", "meal_limit": 200.0, "souvenir_limit": None, "baijiu_limit": 300.0, "red_wine_limit": 300.0},
        {"category": "external", "sub_category": "外事/商务", "personnel_level": "其他人员", "meal_limit": 200.0, "souvenir_limit": 200.0, "baijiu_limit": 200.0, "red_wine_limit": 200.0},
        {"category": "external", "sub_category": "其他公务", "personnel_level": "其他人员", "meal_limit": 150.0, "souvenir_limit": None, "baijiu_limit": 200.0, "red_wine_limit": 200.0},
        # 内部
        {"category": "internal", "sub_category": "", "personnel_level": "省管中层", "meal_limit": 150.0, "souvenir_limit": None, "baijiu_limit": 0.0, "red_wine_limit": 0.0},
        {"category": "internal", "sub_category": "", "personnel_level": "市管中层", "meal_limit": 150.0, "souvenir_limit": None, "baijiu_limit": 0.0, "red_wine_limit": 0.0},
        {"category": "internal", "sub_category": "", "personnel_level": "其他人员", "meal_limit": 100.0, "souvenir_limit": None, "baijiu_limit": 0.0, "red_wine_limit": 0.0},
    ]
    for d in defaults:
        db.add(ReceptionStandard(**d))


def _seed_default_templates(db: Session):
    """从 defaults 注入默认场景模板"""
    from ..standards_query import templates

    for t in templates.get_templates():
        db.add(ReceptionTemplate(
            template_id=t["id"],
            title=t["title"],
            icon=t["icon"],
            icon_class=t["icon_class"],
            reception_type=t["reception_type"],
            category=t["category"],
            example_data=json.dumps(t["example"], ensure_ascii=False),
            notes=t["notes"],
            required_docs=t["required_docs"],
        ))


class OCRRecord(Base):
    """OCR 识别记录 — 关联用户名 + 流水号"""

    __tablename__ = "ocr_records"

    id = Column(Integer, primary_key=True)
    username = Column(String, nullable=False, index=True)       # 操作用户
    serial_number = Column(String, index=True)                  # 关联流水号
    slot_name = Column(String)                                  # 槽位名：审批单/申请单/报账单...
    original_filename = Column(String)                          # 原始文件名
    file_type = Column(String)                                  # image/pdf/doc/xml
    markdown = Column(Text)                                     # OCR 识别的 Markdown 结果
    status = Column(String, default="success")                  # success / error
    error_message = Column(Text)                                # 错误信息
    elapsed = Column(Float)                                     # 耗时(秒)
    created_at = Column(String, default=lambda: datetime.now().isoformat())

    # 坐标定位渲染相关字段
    ocr_blocks = Column(Text)                                   # JSON: 带坐标的 OCR 文字块列表
    ocr_lines = Column(Text)                                    # JSON: OCR 逐行文字列表
    img_width = Column(Integer, default=0)                      # 原始图片宽度
    img_height = Column(Integer, default=0)                     # 原始图片高度
    engine = Column(String)                                     # 识别引擎: rapidocr-gpu / rapidocr-cpu / mineru

    # 原始文件存储
    original_file_path = Column(String)                         # 持久化文件路径（相对于 data/ocr_files/）
    # 去重标识
    dedup_key = Column(String, index=True)                      # slot_name + serial_number + file_hash


class ExtractedField(Base):
    """字段提取结果 — 每个OCR文件提取后的结构化字段"""

    __tablename__ = "extracted_fields"

    id = Column(Integer, primary_key=True)
    serial_number = Column(String, nullable=False, index=True)  # 关联流水号
    slot_name = Column(String, nullable=False, index=True)      # 槽位: 审批单/发票/支付凭证...
    ocr_record_id = Column(Integer)                             # 关联OCR记录ID
    extracted_json = Column(Text)                               # JSON: 提取的结构化字段
    extraction_rule = Column(String)                            # 使用的提取规则名
    status = Column(String, default="pending")                  # pending/extracted/errored
    error_message = Column(Text)                                # 错误信息
    created_at = Column(String, default=lambda: datetime.now().isoformat())
    updated_at = Column(String, default=lambda: datetime.now().isoformat())


class ReviewResult(Base):
    """审核规则校验结果"""

    __tablename__ = "review_results"

    id = Column(Integer, primary_key=True)
    serial_number = Column(String, nullable=False, index=True)
    rule_name = Column(String, nullable=False)                  # 规则名: 金额一致性/日期一致性...
    severity = Column(String, default="高")                     # 高/中/低/提示
    passed = Column(Boolean, default=True)                      # 是否通过
    detail = Column(Text)                                       # 校验详情/不通过原因
    created_at = Column(String, default=lambda: datetime.now().isoformat())


class ExtractionRule(Base):
    """字段提取规则 — 管理员可通过页面配置"""

    __tablename__ = "extraction_rules"

    id = Column(Integer, primary_key=True)
    slot_name = Column(String, nullable=False, unique=True)    # 槽位名: 审批单/申请单/...
    rule_name = Column(String, nullable=False)                  # 规则标识: approval_extraction 等
    fields_json = Column(Text)                                  # JSON: 字段定义列表
    enabled = Column(Boolean, default=True)                     # 是否启用
    # LLM 提取配置
    use_llm = Column(Boolean, default=False)                    # 是否启用大模型提取
    llm_prompt_template = Column(Text)                          # 自定义提示词模板（可选）
    llm_response_schema = Column(Text)                          # JSON Schema，规定返回格式（可选）
    created_at = Column(String, default=lambda: datetime.now().isoformat())
    updated_at = Column(String, default=lambda: datetime.now().isoformat())


class InternalUnit(Base):
    """内部单位通信录 — 通过 Excel 导入/导出管理"""

    __tablename__ = "internal_units"

    id = Column(Integer, primary_key=True)
    name = Column(String, nullable=False, unique=True)  # 公司名称
    status = Column(String, default="active")           # active / deleted
    created_at = Column(String, default=lambda: datetime.now().isoformat())
    updated_at = Column(String, default=lambda: datetime.now().isoformat())

def get_db():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()

def _migrate_internal_units():
    """增量迁移：为已有数据库补加 internal_units 表"""
    import sqlite3
    db_path = DATABASE_URL.replace("sqlite:///", "")
    conn = sqlite3.connect(db_path)
    try:
        cur = conn.cursor()
        cur.execute("SELECT name FROM sqlite_master WHERE type='table' AND name='internal_units'")
        if cur.fetchone() is None:
            cur.execute("""
                CREATE TABLE internal_units (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    name TEXT NOT NULL UNIQUE,
                    status TEXT DEFAULT 'active',
                    created_at TEXT,
                    updated_at TEXT
                )
            """)
            print("Migrated: created internal_units table")
        conn.commit()
    except Exception as e:
        print(f"Migration internal_units failed: {e}")
    finally:
        conn.close()

def _seed_internal_units():
    """确保内部单位数据已注入（首次启动时从 Excel 读取）"""
    db = SessionLocal()
    try:
        if db.query(InternalUnit).first():
            return
        try:
            import openpyxl
        except ImportError:
            print("openpyxl not installed, skipping internal_units seed")
            return
        excel_path = Path(__file__).parent.parent.parent / "审核标准" / "内部单位清单.xlsx"
        if excel_path.exists():
            wb = openpyxl.load_workbook(str(excel_path))
            ws = wb.active
            count = 0
            for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
                if row[0]:
                    unit = InternalUnit(
                        name=str(row[0]).strip(),
                        status="deleted" if row[1] == "删除" else "active",
                    )
                    db.add(unit)
                    count += 1
            db.commit()
            print(f"Seeded internal_units from Excel ({count} rows)")
    except Exception as e:
        db.rollback()
        print(f"Seed internal_units failed: {e}")
    finally:
        db.close()
