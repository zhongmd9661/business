"""从 .docx 制度文档解析规则到结构化数据"""
import json
from pathlib import Path

from .config import RULES_DIR


def _read_docx_text(file_path: Path) -> str:
    """读取 .docx 文本内容"""
    try:
        from docx import Document

        doc = Document(str(file_path))
        texts = [p.text for p in doc.paragraphs if p.text.strip()]
        # Also extract table content
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                texts.append(" | ".join(cells))
        return "\n".join(texts)
    except ImportError:
        # Fallback: read as text if python-docx not available
        return file_path.read_text(encoding="utf-8", errors="ignore")


def _split_into_clauses(text: str) -> list[str]:
    """按条款分段，识别 '第X条' 模式"""
    import re

    # Match patterns like "第一条", "第十二条", etc.
    pattern = re.compile(r"(第[一二三四五六七八九十百千\d]+条)")
    parts = []
    current = ""
    for line in text.split("\n"):
        m = pattern.search(line)
        if m:
            if current:
                parts.append(current)
            current = line
        else:
            current += f"\n{line}"
    if current:
        parts.append(current)
    return parts if len(parts) > 1 else [text]


def _generate_check_expression(category: str, clause_text: str) -> str:
    """根据分类和条款内容生成 check_expression JSON"""
    import re

    # 尝试从文本中提取金额阈值
    amount_match = re.search(r'([\d,]+\.?\d*)\s*元', clause_text)
    threshold = float(amount_match.group(1).replace(',', '')) if amount_match else None

    # 尝试提取日期
    date_match = re.search(r'(\d{4}[-/]\d{2}[-/]\d{2})', clause_text)
    date_val = date_match.group(1).replace('/', '-') if date_match else None

    expr = None

    if category == "金额标准":
        # 金额比较规则
        if threshold:
            expr = {
                "type": "amount_compare",
                "field": "per_person_amount",
                "operator": ">",
                "threshold": threshold,
            }
        else:
            expr = {
                "type": "keyword_match",
                "keywords": ["超标", "超过", "超出"],
                "mode": "any",
            }

    elif category == "禁止性规定":
        # 关键词匹配 — 从条款文本中提取禁止性关键词
        keywords = []
        for kw in ("私人会所", "高档娱乐", "一桌餐", "鱼翅", "燕窝",
                    "烟", "酒", "送礼", "现金", "购物卡", "旅游", "景点"):
            if kw in clause_text:
                keywords.append(kw)
        if keywords:
            expr = {
                "type": "keyword_match",
                "keywords": keywords,
                "mode": "any",
            }
        else:
            expr = {
                "type": "keyword_match",
                "keywords": ["禁止", "不得", "严禁"],
                "mode": "any",
            }

    elif category == "单据完整性":
        # 布尔检查 — 检查必需单据是否存在
        if any(k in clause_text for k in ["支付凭证", "刷卡单"]):
            expr = {"type": "boolean_check", "field": "payment_voucher_present", "expected": False}
        elif any(k in clause_text for k in ["支付流水", "交易流水"]):
            expr = {"type": "boolean_check", "field": "payment_statement_present", "expected": False}
        elif any(k in clause_text for k in ["往来公函", "公函"]):
            expr = {"type": "boolean_check", "field": "official_letter_present", "expected": False}
        elif any(k in clause_text for k in ["发票查验", "发票验证"]):
            expr = {"type": "boolean_check", "field": "invoice_verified", "expected": False}
        elif any(k in clause_text for k in ["费用明细", "明细清单"]):
            expr = {"type": "boolean_check", "field": "expense_detail_list_present", "expected": False}
        elif any(k in clause_text for k in ["网格分摊", "分摊表"]):
            expr = {"type": "boolean_check", "field": "grid_allocation_signed", "expected": False}
        else:
            expr = {
                "type": "keyword_match",
                "keywords": ["缺少", "缺失", "未提供"],
                "mode": "any",
            }

    elif category in ("审批管理", "招待类型"):
        expr = {
            "type": "keyword_match",
            "keywords": ["审批", "事前", "备案", "报备"],
            "mode": "any",
        }

    elif category == "陪同人数":
        if "内部" in clause_text:
            expr = {
                "type": "count_compare",
                "field": "companion_count",
                "compare_field": "guest_count",
                "operator": ">",
                "formula": "fixed_if_le_then_ratio",
                "fixed_limit": 3,
                "ratio": 3,
            }
        else:
            expr = {
                "type": "count_compare",
                "field": "companion_count",
                "compare_field": "guest_count",
                "operator": ">",
                "formula": "equal_if_le_5",
            }

    elif category == "报销合规":
        if any(k in clause_text for k in ["现金", "现金支付"]):
            expr = {
                "type": "boolean_check",
                "field": "is_cash_payment",
                "expected": True,
                "condition": {"actual_amount": 5000},
            }
        elif any(k in clause_text for k in ["预存", "签单"]):
            expr = {"type": "boolean_check", "field": "has_prepaid", "expected": True}
        else:
            expr = {
                "type": "keyword_match",
                "keywords": ["拆分", "混淆", "专票"],
                "mode": "any",
            }

    elif category == "交叉稽核":
        expr = {
            "type": "keyword_match",
            "keywords": ["重复", "虚假", "差旅", "风险"],
            "mode": "any",
        }

    if expr is None:
        # 兜底：关键词匹配
        expr = {
            "type": "keyword_match",
            "keywords": ["违规", "问题", "风险"],
            "mode": "any",
        }

    return json.dumps(expr, ensure_ascii=False)


def _classify_clause(clause_text: str, source: str) -> dict:
    """根据条款内容分类并提取规则信息"""
    import re

    m = re.search(r"第[一二三四五六七八九十百千\d]+条", clause_text)
    clause_ref = m.group(0) if m else ""

    # Default classification
    category = "单据完整性"
    level = "中"
    rule_name = clause_ref

    # Classify by keywords
    if any(k in clause_text for k in ["金额", "标准", "元/人", "不得超过"]):
        category = "金额标准"
        level = "高"
    elif any(k in clause_text for k in ["招待类型", "商务招待", "外事招待", "内部业务招待", "其他公务"]):
        category = "招待类型"
        level = "中"
    elif any(k in clause_text for k in ["陪同人数", "陪餐人数", "陪同"]):
        category = "陪同人数"
        level = "高"
    elif any(k in clause_text for k in ["严禁", "禁止", "不得", "严禁", "不讲排场"]):
        category = "禁止性规定"
        level = "高"
    elif any(k in clause_text for k in ["报销", "发票", "支付凭证", "结算", "报账"]):
        category = "报销合规"
        level = "中"
    elif any(k in clause_text for k in ["审批", "事前", "备案", "报备"]):
        category = "审批管理"
        level = "中"
    elif any(k in clause_text for k in ["交叉", "重复", "差旅", "风险", "虚假"]):
        category = "交叉稽核"
        level = "高"

    # Extract rule name from first meaningful line
    lines = [l.strip() for l in clause_text.split("\n") if l.strip()]
    if lines:
        rule_name = lines[0][:50]

    check_expression = _generate_check_expression(category, clause_text)

    return {
        "category": category,
        "rule_name": rule_name,
        "clause": f"{source} {clause_ref}",
        "level": level,
        "description": clause_text[:500],
        "check_expression": check_expression,
        "source_document": source,
    }


def parse_document(file_path: Path) -> list[dict]:
    """解析单个制度文档"""
    text = _read_docx_text(file_path)
    source = file_path.stem
    clauses = _split_into_clauses(text)
    rules = []
    for clause in clauses:
        if len(clause.strip()) > 10:
            rule = _classify_clause(clause, source)
            rules.append(rule)
    return rules


def parse_all_documents(use_llm: bool = False) -> list[dict]:
    """解析所有制度文档

    Args:
        use_llm: 是否使用 LLM 辅助解析（默认用关键词分类）

    Returns:
        结构化规则列表
    """
    if not RULES_DIR.exists():
        return []

    if use_llm:
        return _parse_all_with_llm()

    all_rules = []
    for fpath in sorted(RULES_DIR.glob("*.docx")):
        rules = parse_document(fpath)
        all_rules.extend(rules)
    return all_rules


def _parse_all_with_llm() -> list[dict]:
    """使用 LLM 辅助解析所有制度文档"""
    from .rule_llm_parser import LlmRuleParser

    clauses: list[tuple[str, str]] = []
    for fpath in sorted(RULES_DIR.glob("*.docx")):
        text = _read_docx_text(fpath)
        source = fpath.stem
        split_clauses = _split_into_clauses(text)
        for clause in split_clauses:
            if len(clause.strip()) > 10:
                clauses.append((source, clause))

    if not clauses:
        return []

    parser = LlmRuleParser()
    try:
        parser.start()
        return parser.parse_clauses(clauses)
    except Exception:
        # LLM 不可用时回退到关键词分类
        all_rules = []
        for source, clause in clauses:
            rule = _classify_clause(clause, source)
            all_rules.append(rule)
        return all_rules
    finally:
        parser.stop()
