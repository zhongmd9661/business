"""内部单位通信录管理 API — CRUD + Excel 导入/导出/模板下载"""

import io
import json
from datetime import datetime

from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from sqlalchemy.orm import Session

from .auth import get_admin_user, get_current_user
from .models_db import InternalUnit, User, get_db

router = APIRouter()


# ===================================================================
# 注意：具体路由（import/export/template）必须放在 {unit_id} 参数路由之前
# ===================================================================

@router.post("/admin/internal-units/import", summary="Excel 批量导入")
def import_internal_units(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    admin: User = Depends(get_admin_user),
):
    """通过 Excel 文件批量导入内部单位"""
    import openpyxl

    if not file.filename.endswith((".xlsx", ".xls")):
        raise HTTPException(400, "仅支持 .xlsx 或 .xls 文件")

    content = file.file.read()
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content))
        ws = wb.active
    except Exception as e:
        raise HTTPException(400, f"Excel 解析失败: {e}")

    added = 0
    updated = 0
    deleted = 0
    ignored = 0

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
        company_name = row[0]
        status_val = row[1] if len(row) > 1 else "增加"
        if not company_name or not str(company_name).strip():
            continue
        company_name = str(company_name).strip()
        status_val = str(status_val).strip() if status_val else "增加"

        if status_val == "删除":
            existing = db.query(InternalUnit).filter(InternalUnit.name == company_name).first()
            if existing:
                existing.status = "deleted"
                existing.updated_at = datetime.now().isoformat()
                deleted += 1
            else:
                ignored += 1
        else:
            existing = db.query(InternalUnit).filter(InternalUnit.name == company_name).first()
            if existing:
                existing.status = "active"
                existing.updated_at = datetime.now().isoformat()
                updated += 1
            else:
                db.add(InternalUnit(name=company_name, status="active"))
                added += 1

    db.commit()
    return {
        "total_rows": ws.max_row - 1,
        "added": added,
        "updated": updated,
        "deleted": deleted,
        "ignored": ignored,
    }


@router.get("/admin/internal-units/export", summary="导出内部单位为 Excel")
def export_internal_units(db: Session = Depends(get_db), _admin: User = Depends(get_admin_user)):
    """导出所有内部单位为 Excel 文件"""
    units = db.query(InternalUnit).order_by(InternalUnit.name).all()
    wb = Workbook()
    ws = wb.active
    ws.title = "内部单位清单"
    ws.append(["公司名称", "状态"])
    for u in units:
        ws.append([u.name, "删除" if u.status == "deleted" else "增加"])
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 10

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=internal_units.xlsx"},
    )


@router.get("/admin/internal-units/template", summary="下载导入模板")
def download_template():
    """下载 Excel 导入模板"""
    wb = Workbook()
    ws = wb.active
    ws.title = "内部单位清单"
    ws.append(["公司名称", "状态"])
    ws.append(["示例公司A", "增加"])
    ws.append(["示例公司B", "删除"])
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 10

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=内部单位导入模板.xlsx"},
    )


# ===================================================================
# 通用 CRUD（放最后，避免 {unit_id} 匹配到 import/export/template）
# ===================================================================

@router.get("/admin/internal-units", summary="获取内部单位列表")
def list_internal_units(
    status: str | None = None,
    db: Session = Depends(get_db),
    _admin: User = Depends(get_admin_user),
):
    """获取所有内部单位（可按状态筛选）"""
    query = db.query(InternalUnit)
    if status:
        query = query.filter(InternalUnit.status == status)
    units = query.order_by(InternalUnit.name).all()
    return [_unit_to_dict(u) for u in units]


@router.post("/admin/internal-units", summary="新增内部单位")
def create_internal_unit(
    data: dict,
    db: Session = Depends(get_db),
    admin: User = Depends(get_admin_user),
):
    """新增一个内部单位"""
    name = data.get("name", "").strip()
    if not name:
        raise HTTPException(400, "公司名称不能为空")
    existing = db.query(InternalUnit).filter(InternalUnit.name == name).first()
    if existing:
        raise HTTPException(400, f"单位「{name}」已存在")
    unit = InternalUnit(name=name, status=data.get("status", "active"))
    db.add(unit)
    db.commit()
    db.refresh(unit)
    return _unit_to_dict(unit)


@router.put("/admin/internal-units/{unit_id}", summary="修改内部单位状态")
def update_internal_unit(
    unit_id: int,
    data: dict,
    db: Session = Depends(get_db),
    admin: User = Depends(get_admin_user),
):
    """修改内部单位名称或状态"""
    unit = db.query(InternalUnit).filter(InternalUnit.id == unit_id).first()
    if not unit:
        raise HTTPException(404, "单位不存在")
    if "name" in data:
        new_name = data["name"].strip()
        if not new_name:
            raise HTTPException(400, "公司名称不能为空")
        dup = db.query(InternalUnit).filter(InternalUnit.name == new_name, InternalUnit.id != unit_id).first()
        if dup:
            raise HTTPException(400, f"单位「{new_name}」已存在")
        unit.name = new_name
    if "status" in data:
        unit.status = data["status"]
    unit.updated_at = datetime.now().isoformat()
    db.commit()
    db.refresh(unit)
    return _unit_to_dict(unit)


@router.delete("/admin/internal-units/{unit_id}", summary="删除内部单位")
def delete_internal_unit(
    unit_id: int,
    db: Session = Depends(get_db),
    admin: User = Depends(get_admin_user),
):
    """删除内部单位（硬删除）"""
    unit = db.query(InternalUnit).filter(InternalUnit.id == unit_id).first()
    if not unit:
        raise HTTPException(404, "单位不存在")
    db.delete(unit)
    db.commit()
    return {"detail": "已删除", "id": unit_id}


def _unit_to_dict(u: InternalUnit) -> dict:
    return {
        "id": u.id,
        "name": u.name,
        "status": u.status,
        "created_at": u.created_at,
        "updated_at": u.updated_at,
    }
