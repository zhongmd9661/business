"""专业报表导出引擎 - 将重建后的虚拟大表还原为具备滚动特性的 Excel 文件"""
from __future__ import annotations

import pathlib
from typing import List, Any
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.drawing.image import Image as OpenpyxlImage

class ExcelReporter:
    """实现具备报表系统特性的 Excel 导出"""

    def generate(self, table_data: List[List[Any]], headers: List[str], output_path: pathlib.Path, image_mappings: Dict[str, pathlib.Path] = None):
        """
        生成高保真 Excel 报表
        table_data: [Row][Col] 数据矩阵
        headers: 全局列名
        image_mappings: { row_id: image_path } 用于在表中嵌入原图对照 (可选)
        """
        wb = Workbook()
        ws = wb.active
        ws.title = "重建报表"

        # 1. 写入表头
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = Font(bold=True, size=12)
            cell.alignment = Alignment(horizontal="center")
            # 添加边框
            thin = Side(border_style="thin", color="000000")
            cell.border = Border(top=thin, left=thin, right=thin, bottom=thin)

        # 2. 写入数据行
        for row_idx, row_values in enumerate(table_data, 2):
            for col_idx, value in enumerate(row_values, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                thin = Side(border_style="thin", color="000000")
                cell.border = Border(top=thin, left=thin, right=thin, bottom=thin)

        # 3. 实现报表系统滚动特性 (冻结窗格)
        # 冻结首行和首列
        ws.freeze_panes = "B2"

        # 4. 列宽自动调整 (简单实现)
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except: pass
            ws.column_dimensions[column].width = max_length + 2

        # 5. (高级) 图片嵌入 - 如果提供了映射且需要原图对照
        if image_mappings:
            for row_idx, img_path in image_mappings.items():
                if pathlib.Path(img_path).exists():
                    img = OpenpyxlImage(img_path)
                    # 适当缩放图片以适应单元格 (示例：放在数据右侧或单独 Sheet)
                    # 为简化演示，此处暂不实现复杂嵌入，仅保留接口
                    pass

        wb.save(output_path)
